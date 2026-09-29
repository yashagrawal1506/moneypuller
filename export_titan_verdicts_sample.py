"""Export the Titan verdict sample to Excel (any-target vs SL simulation)."""

import sqlite3
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DB = Path("data/moneypuller.db")
OUT = Path("data/titan_verdicts_sample.xlsx")

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
FILL_OPEN = PatternFill("solid", fgColor="E7E6E6")
FILL_SUM = PatternFill("solid", fgColor="DDEBF7")
PRICE_FMT = "#,##0.00"
PCT_FMT = "+0.0%;-0.0%"

SQL = """
SELECT a.record_date, r.stock_name, r.action, r.cmp, v.entry_price,
       r.target_1, r.target_2, r.target_3, r.stop_loss,
       v.status, v.exit_price, v.exit_level, v.exit_date,
       v.days_taken, v.pct, v.note
FROM verdicts v
JOIN recommendations r ON r.id = v.rec_id
JOIN articles a ON a.id = r.article_id
WHERE r.stock_name = 'Titan Company'
ORDER BY a.record_date
"""

HEADERS = ["Record Date", "Stock", "Action", "Article CMP", "Entry (Open)",
           "T1", "T2", "T3", "SL", "Achieved Target or SL?", "Exit Price",
           "Exit Level", "Exit Date", "Days Taken", "%age", "Note"]
WIDTHS = [12, 15, 8, 12, 12, 9, 9, 9, 9, 20, 11, 12, 12, 10, 9, 46]


def main() -> None:
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    rows = conn.execute(SQL).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Titan Verdicts"

    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    ws.row_dimensions[1].height = 30
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for r, row in enumerate(rows, start=2):
        vals = [row["record_date"], row["stock_name"], row["action"],
                row["cmp"], row["entry_price"], row["target_1"],
                row["target_2"], row["target_3"], row["stop_loss"],
                row["status"], row["exit_price"], row["exit_level"],
                row["exit_date"], row["days_taken"], row["pct"], row["note"]]
        for j, val in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=val)
            c.border = BORDER
            if j in (4, 5, 6, 7, 8, 9, 11):
                c.number_format = PRICE_FMT
            if j == 15:
                c.number_format = PCT_FMT
            if j == 10:
                status = row["status"]
                if status == "target achieved":
                    c.fill = FILL_OK
                elif status == "SL achieved":
                    c.fill = FILL_BAD
                else:
                    c.fill = FILL_OPEN

    # summary block
    n = len(rows)
    srow = n + 4
    n_t = sum(1 for x in rows if x["status"] == "target achieved")
    n_s = sum(1 for x in rows if x["status"] == "SL achieved")
    n_a = sum(1 for x in rows if x["status"].startswith("AMBIGUOUS"))
    n_o = sum(1 for x in rows if x["status"] == "NO HIT YET")
    wins = [x["pct"] for x in rows if x["status"] == "target achieved"
            and x["pct"] is not None]
    losses = [x["pct"] for x in rows if x["status"] == "SL achieved"
              and x["pct"] is not None]
    live = [x["pct"] for x in rows if x["status"] == "NO HIT YET"
            and x["pct"] is not None]
    summary = [
        ("Evaluated", n),
        ("Target achieved", n_t),
        ("SL achieved", n_s),
        ("AMBIGUOUS", n_a),
        ("NO HIT YET (live)", n_o),
        ("Win rate (of settled)", f"{n_t}/{n_t + n_s} = {n_t / (n_t + n_s):.0%}"),
        ("Avg % on targets", f"{sum(wins) / len(wins):+.1%}" if wins else "-"),
        ("Avg % on SL", f"{sum(losses) / len(losses):+.1%}" if losses else "-"),
        ("Avg % all settled", f"{(sum(wins) + sum(losses)) / (len(wins) + len(losses)):+.1%}"
         if wins or losses else "-"),
        ("Live P&L (unrealised)", f"{sum(live):+.1%}" if live else "-"),
        ("Rule", "Entry = record-date OPEN; exit = furthest target hit before "
                 "SL (else SL); % = exit/entry-1 (Buy), (entry-exit)/entry (Sell); "
                 "days = trading sessions incl. exit day"),
    ]
    for i, (k, v) in enumerate(summary, start=srow):
        kc = ws.cell(row=i, column=1, value=k)
        vc = ws.cell(row=i, column=2, value=v)
        kc.font = Font(bold=True)
        kc.fill = FILL_SUM
        vc.fill = FILL_SUM

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:P{n + 1}"
    wb.save(OUT)
    print(f"wrote {OUT} ({n} leads)")


if __name__ == "__main__":
    main()
