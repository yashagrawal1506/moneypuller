"""
MoneyPuller - export the leads database to a formatted Excel workbook.

Reads ONLY from data/moneypuller.db (no network). Produces:

  Sheet "Leads"     - one row per recommendation with the verdict simulation
                      columns (Achieved Target or SL?, entry/exit, %age,
                      days taken), derived upside/risk formulas, colored cells
  Sheet "Summary"   - KPIs, action breakdown, per-analyst and per-date tables
"""

import sqlite3
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEFAULT_XLSX = Path("data/moneypuller_leads.xlsx")

# ---------- styling ----------
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
KPI_FONT = Font(bold=True)
FILL_BUY = PatternFill("solid", fgColor="C6EFCE")
FILL_AVOID = PatternFill("solid", fgColor="FFEB9C")
FILL_SELL = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

PRICE_FMT = "#,##0.00"
PCT_FMT = "0.0%"
DATE_FMT = "yyyy-mm-dd"

# (header, sql key, width, number format)
COLUMNS = [
    ("Record Date", "record_date", 12, DATE_FMT),
    ("Stock Name", "stock_name", 30, None),
    ("CMP", "cmp", 11, PRICE_FMT),
    ("Action", "action", 9, None),
    ("Target 1", "target_1", 10, PRICE_FMT),
    ("Target 2", "target_2", 10, PRICE_FMT),
    ("Target 3", "target_3", 10, PRICE_FMT),
    ("Stop-Loss", "stop_loss", 10, PRICE_FMT),
    ("Achieved Target or SL?", "verdict_status", 22, None),
    ("Entry (Open)", "entry_price", 11, PRICE_FMT),
    ("Exit Price", "exit_price", 11, PRICE_FMT),
    ("Exit Level", "exit_level", 16, None),
    ("Exit Date", "exit_date", 12, DATE_FMT),
    ("Days Taken", "days_taken", 10, "0"),
    ("%age", "pct", 9, PCT_FMT),
    ("Confidence %", "confidence", 12, PCT_FMT),
    ("Verdict Note", "verdict_note", 42, None),
    ("T1 Upside %", None, 11, PCT_FMT),
    ("T2 Upside %", None, 11, PCT_FMT),
    ("T3 Upside %", None, 11, PCT_FMT),
    ("Risk % (SL)", None, 10, PCT_FMT),
    ("R:R (T1)", None, 9, "0.0"),
    ("Analyst Name", "analyst_name", 20, None),
    ("Analyst (byline)", "analyst", 36, None),
    ("Reasoning", "reasoning", 85, None),
    ("Article Title", "article_title", 50, None),
    ("Article URL", "article_url", 45, None),
]
# 1-based sheet column numbers for formula references
COL = {h: i + 1 for i, (h, _, _, _) in enumerate(COLUMNS)}
L = {h: get_column_letter(n) for h, n in COL.items()}


def _fetch_rows(conn: sqlite3.Connection) -> list[sqlite3.Row]:
    cur = conn.execute(
        """SELECT a.record_date, a.title AS article_title, a.url AS article_url,
                  r.stock_name, r.cmp, r.action, r.target_1, r.target_2,
                  r.target_3, r.stop_loss, r.reasoning, r.analyst,
                  v.status AS verdict_status, v.entry_price, v.exit_price,
                  v.exit_level, v.exit_date, v.days_taken, v.pct,
                  v.confidence, v.note AS verdict_note
           FROM recommendations r
           JOIN articles a ON a.id = r.article_id
           LEFT JOIN verdicts v ON v.rec_id = r.id
           ORDER BY CASE WHEN a.record_date IS NULL THEN 1 ELSE 0 END,
                    a.record_date DESC, r.id""")
    return cur.fetchall()


def _analyst_name(raw: str | None) -> str | None:
    """Canonical analyst (person) name: the byline's first comma segment.

    Moneycontrol formats the same analyst differently (title, punctuation,
    firm rebrands and even job changes over time), but the name before the
    first comma is stable - e.g. all 8 'Amol Athawale, ...' variants collapse
    to 'Amol Athawale'. Verified: no two distinct people share a name in
    this dataset, so name-only grouping is safe for pivots.
    """
    if not raw or not raw.strip():
        return None
    return raw.split(",")[0].strip() or None


def _action_fill(action: str | None) -> PatternFill | None:
    a = (action or "").lower()
    if "buy" in a:
        return FILL_BUY
    if "avoid" in a or "hold" in a:
        return FILL_AVOID
    if "sell" in a:
        return FILL_SELL
    return None


def _iso_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        return None


def _write_leads_sheet(wb: Workbook, rows: list[sqlite3.Row]) -> None:
    ws = wb.active
    ws.title = "Leads"

    for i, (header, _, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22

    for r, row in enumerate(rows, start=2):
        d = dict(row)
        rec_date = _iso_date(d["record_date"])
        values = {
            "record_date": rec_date, "stock_name": d["stock_name"],
            "cmp": d["cmp"], "action": d["action"],
            "target_1": d["target_1"], "target_2": d["target_2"],
            "target_3": d["target_3"], "stop_loss": d["stop_loss"],
            "verdict_status": d["verdict_status"],
            "entry_price": d["entry_price"], "exit_price": d["exit_price"],
            "exit_level": d["exit_level"],
            "exit_date": _iso_date(d["exit_date"]),
            "days_taken": d["days_taken"], "pct": d["pct"],
            "confidence": d["confidence"],
            "verdict_note": d["verdict_note"],
            "analyst_name": _analyst_name(d["analyst"]),
            "analyst": d["analyst"], "reasoning": d["reasoning"],
            "article_title": d["article_title"], "article_url": d["article_url"],
        }
        for i, (header, key, _, numfmt) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=i)
            if key is not None:
                cell.value = values.get(key)
            elif header == "T1 Upside %":
                cell.value = (f'=IF(OR(${L["CMP"]}{r}="",${L["Target 1"]}{r}=""),"",'
                              f'${L["Target 1"]}{r}/${L["CMP"]}{r}-1)')
            elif header == "T2 Upside %":
                cell.value = (f'=IF(OR(${L["CMP"]}{r}="",${L["Target 2"]}{r}=""),"",'
                              f'${L["Target 2"]}{r}/${L["CMP"]}{r}-1)')
            elif header == "T3 Upside %":
                cell.value = (f'=IF(OR(${L["CMP"]}{r}="",${L["Target 3"]}{r}=""),"",'
                              f'${L["Target 3"]}{r}/${L["CMP"]}{r}-1)')
            elif header == "Risk % (SL)":
                cell.value = (f'=IF(OR(${L["CMP"]}{r}="",${L["Stop-Loss"]}{r}=""),"",'
                              f'${L["Stop-Loss"]}{r}/${L["CMP"]}{r}-1)')
            elif header == "R:R (T1)":
                up, risk = f'{L["T1 Upside %"]}{r}', f'{L["Risk % (SL)"]}{r}'
                cell.value = f'=IFERROR({up}/ABS({risk}),"")'
            if numfmt and cell.value is not None:
                cell.number_format = numfmt
            cell.border = BORDER
            cell.alignment = WRAP_TOP if header == "Reasoning" else TOP

        fill = _action_fill(d["action"])
        if fill:
            ws.cell(row=r, column=COL["Action"]).fill = fill
        status_fill = {"target achieved": FILL_BUY,
                       "SL achieved": FILL_SELL,
                       "AMBIGUOUS (target & SL same day)": FILL_AVOID,
                       }.get(d["verdict_status"] or "")
        if status_fill:
            ws.cell(row=r, column=COL["Achieved Target or SL?"]).fill = status_fill
        # keep wrapped reasoning rows readable
        reasoning = d["reasoning"] or ""
        if reasoning:
            lines = -(-len(reasoning) // 95)  # ceil
            ws.row_dimensions[r].height = min(15 * max(lines, 1) + 4, 140)

    ws.freeze_panes = "C2"
    last = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last}{max(len(rows) + 1, 2)}"


def _write_summary_sheet(wb: Workbook, rows: list[sqlite3.Row], db_path: Path) -> None:
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    for col in "BCDE":
        ws.column_dimensions[col].width = 16

    ws["A1"] = "MoneyPuller - Trade Spotlight Leads"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now():%Y-%m-%d %H:%M} from {db_path} (snapshot, not live)"

    dates = sorted({d for r in rows if (d := _iso_date(r["record_date"]))})
    cmps = [r["cmp"] for r in rows if r["cmp"]]
    kpis = [
        ("Articles", len({r["article_url"] for r in rows})),
        ("Leads", len(rows)),
        ("Unique stocks", len({r["stock_name"] for r in rows})),
        ("Unique analysts", len({r["analyst"] for r in rows if r["analyst"]})),
        ("Date range", f"{dates[0]:%d %b %Y} - {dates[-1]:%d %b %Y}" if dates else "-"),
        ("Avg CMP of leads", round(sum(cmps) / len(cmps), 2) if cmps else "-"),
    ]
    ws["A4"] = "Key figures"
    ws["A4"].font = KPI_FONT
    for i, (k, v) in enumerate(kpis, start=5):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)

    # verdict outcome KPIs (only dated leads can be evaluated)
    settled = [r for r in rows if r["verdict_status"] in
               ("target achieved", "SL achieved")]
    n_tgt = sum(1 for r in settled if r["verdict_status"] == "target achieved")
    n_sl = len(settled) - n_tgt
    pcts = [r["pct"] for r in settled if r["pct"] is not None]
    tgt_pcts = [r["pct"] for r in settled
                if r["verdict_status"] == "target achieved"
                and r["pct"] is not None]
    sl_pcts = [r["pct"] for r in settled if r["verdict_status"] == "SL achieved"
               and r["pct"] is not None]
    days = [r["days_taken"] for r in settled if r["days_taken"] is not None]
    vstatus = Counter(r["verdict_status"] or "not evaluated" for r in rows)
    vkpis = [
        ("Verdicts: target achieved", n_tgt),
        ("Verdicts: SL achieved", n_sl),
        ("Verdicts: AMBIGUOUS", vstatus.get("AMBIGUOUS (target & SL same day)", 0)),
        ("Verdicts: NO HIT YET (live)", vstatus.get("NO HIT YET", 0)),
        ("Verdicts: no price data", vstatus.get("NO PRICE DATA", 0)
         + vstatus.get("NOT EVALUABLE (no targets, no SL)", 0)),
        ("Win rate (settled)", f"{n_tgt}/{len(settled)} = "
         f"{n_tgt / len(settled):.0%}" if settled else "-"),
        ("Avg % on target exits", round(sum(tgt_pcts) / len(tgt_pcts), 4)
         if tgt_pcts else "-"),
        ("Avg % on SL exits", round(sum(sl_pcts) / len(sl_pcts), 4)
         if sl_pcts else "-"),
        ("Avg % all settled", round(sum(pcts) / len(pcts), 4) if pcts else "-"),
        ("Avg trading days to exit", round(sum(days) / len(days), 1)
         if days else "-"),
    ]
    live_conf = [r["confidence"] for r in rows
                 if r["verdict_status"] == "NO HIT YET"
                 and r["confidence"] is not None]
    if live_conf:
        vkpis.append(("Avg confidence of live leads",
                      round(sum(live_conf) / len(live_conf), 4)))


    actions = Counter((r["action"] or "Unknown").title() for r in rows)
    ws["A12"] = "Action breakdown"
    ws["A12"].font = KPI_FONT
    ws["A13"], ws["B13"] = "Action", "Leads"
    for cell in (ws["A13"], ws["B13"]):
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BORDER
    for i, (action, n) in enumerate(sorted(actions.items(), key=lambda x: -x[1]), start=14):
        ws.cell(row=i, column=1, value=action).border = BORDER
        c = ws.cell(row=i, column=2, value=n)
        c.border = BORDER
        if fill := _action_fill(action):
            c.fill = fill

    act_end = 14 + max(len(actions), 1) - 1
    vk_row = act_end + 2
    ws.cell(row=vk_row, column=1, value="Verdict simulation "
            "(entry = record-date OPEN, exit = highest target hit, else SL)")
    ws.cell(row=vk_row, column=1).font = KPI_FONT
    for i, (k, v) in enumerate(vkpis, start=vk_row + 1):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, float):
            c.number_format = PCT_FMT if "Avg %" in k else "0.0"

    start = vk_row + len(vkpis) + 2
    ws.cell(row=start, column=1, value="Leads by analyst (canonical name)").font = KPI_FONT
    head = ["Analyst Name", "Leads", "Buys", "Avg T1 upside %", "Avg risk % (SL)"]
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=start + 1, column=j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
    by_analyst: dict[str, list] = {}
    for r in rows:
        by_analyst.setdefault(_analyst_name(r["analyst"]) or "Unknown", []).append(r)
    for i, (analyst, rs) in enumerate(
            sorted(by_analyst.items(), key=lambda x: (-len(x[1]), x[0])), start=start + 2):
        t1s = [r["target_1"] / r["cmp"] - 1 for r in rs
               if r["target_1"] and r["cmp"]]
        risks = [r["stop_loss"] / r["cmp"] - 1 for r in rs
                 if r["stop_loss"] and r["cmp"]]
        ws.cell(row=i, column=1, value=analyst).border = BORDER
        ws.cell(row=i, column=2, value=len(rs)).border = BORDER
        ws.cell(row=i, column=3, value=sum(1 for r in rs if "buy" in (r["action"] or "").lower())).border = BORDER
        c = ws.cell(row=i, column=4, value=round(sum(t1s) / len(t1s), 4) if t1s else None)
        c.number_format = PCT_FMT
        c.border = BORDER
        c = ws.cell(row=i, column=5, value=round(sum(risks) / len(risks), 4) if risks else None)
        c.number_format = PCT_FMT
        c.border = BORDER

    start2 = i + 3 if rows else start + 4
    ws.cell(row=start2, column=1, value="Leads by record date").font = KPI_FONT
    for j, h in enumerate(["Date", "Leads", "Buys"], start=1):
        c = ws.cell(row=start2 + 1, column=j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
    by_date: dict[str, list] = {}
    for r in rows:
        by_date.setdefault(r["record_date"] or "Unknown", []).append(r)
    for i, (d, rs) in enumerate(sorted(by_date.items(), reverse=True), start=start2 + 2):
        ws.cell(row=i, column=1, value=_iso_date(d) if d != "Unknown" else d).border = BORDER
        ws.cell(row=i, column=1).number_format = DATE_FMT
        ws.cell(row=i, column=2, value=len(rs)).border = BORDER
        ws.cell(row=i, column=3, value=sum(1 for r in rs if "buy" in (r["action"] or "").lower())).border = BORDER


def export_workbook(out_path: Path = DEFAULT_XLSX,
                    db_path: Path | None = None) -> int:
    """Export the current DB contents to an xlsx file. Returns lead count."""
    from moneypuller.db import DB_PATH, connect

    db_path = db_path or DB_PATH
    if not Path(db_path).exists():
        raise FileNotFoundError(f"Database not found at {db_path} - run the scraper first")

    conn = connect(db_path)
    try:
        rows = _fetch_rows(conn)
    finally:
        conn.close()

    wb = Workbook()
    _write_leads_sheet(wb, rows)
    _write_summary_sheet(wb, rows, Path(db_path))
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return len(rows)
