"""Export the price database to two Excel workbooks (FY split, per request).

  File 1: data/MoneyPuller_Prices_FY2025-26.xlsx  (2025-04-01 -> 2026-03-31)
  File 2: data/MoneyPuller_Prices_FY2026-27.xlsx  (2026-04-01 -> latest)

Each workbook contains:
  Prices          - date, symbol, OHLC, adjclose, volume (bulk sheet)
  Leads           - recommendations whose record_date falls in the FY
  Dividends_Splits- corporate actions in the FY
  Symbol_Meta     - per-symbol metadata (name, 52wk range, sessions)
  Info            - generation notes and known gaps
"""

import sqlite3
import sys
import time
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DB = Path("data/moneypuller.db")
PRICE_FMT = "#,##0.00"
VOL_FMT = "#,##0"
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True)

FY1 = ("2025-04-01", "2026-03-31")
FY2 = ("2026-04-01", "2099-12-31")


def header_row(ws, headers, widths):
    cells = []
    for h in headers:
        c = WriteOnlyCell(ws, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        cells.append(c)
    ws.append(cells)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build(conn, fy: tuple[str, str], out: Path, label: str) -> dict:
    lo, hi = fy
    t0 = time.time()
    wb = Workbook(write_only=True)

    # ---- Prices (bulk) ----
    ws = wb.create_sheet("Prices")
    header_row(ws, ["Date", "Symbol", "Open", "High", "Low", "Close",
                    "Adj Close", "Volume"],
               [12, 14, 11, 11, 11, 11, 11, 13])
    n_prices = 0
    cur = conn.execute(
        """SELECT date, symbol, open, high, low, close, adjclose, volume
           FROM prices WHERE date BETWEEN ? AND ? ORDER BY date, symbol""",
        (lo, hi))
    for d, sym, o, h, l, c, ac, v in cur:
        row = [d, sym, o, h, l, c, ac, v]
        if n_prices < 2 or v == 0:
            styled = []
            for i, val in enumerate(row):
                cell = WriteOnlyCell(ws, value=val)
                if i >= 2:
                    cell.number_format = VOL_FMT if i == 7 else PRICE_FMT
                styled.append(cell)
            ws.append(styled)
        else:
            ws.append(row)
        n_prices += 1

    # ---- Leads in this FY ----
    ws2 = wb.create_sheet("Leads")
    header_row(ws2, ["Record Date", "Stock", "Yahoo Symbol", "Action", "CMP",
                     "T1", "T2", "T3", "Stop-Loss", "Analyst"],
               [12, 34, 14, 8, 11, 10, 10, 10, 10, 42])
    n_leads = 0
    for row in conn.execute(
            """SELECT a.record_date, r.stock_name, s.yahoo_symbol, r.action,
                      r.cmp, r.target_1, r.target_2, r.target_3, r.stop_loss,
                      r.analyst
               FROM recommendations r
               JOIN articles a ON a.id = r.article_id
               LEFT JOIN symbols s ON s.stock_name = r.stock_name
               WHERE a.record_date BETWEEN ? AND ?
               ORDER BY a.record_date, r.stock_name""", (lo, hi)):
        ws2.append(row)
        n_leads += 1

    # ---- Dividends & splits ----
    ws3 = wb.create_sheet("Dividends_Splits")
    header_row(ws3, ["Date", "Symbol", "Type", "Amount / Ratio"],
               [12, 14, 10, 16])
    n_events = 0
    for row in conn.execute(
            """SELECT date, symbol, 'DIV', amount FROM dividends
               WHERE date BETWEEN ? AND ?
               UNION ALL
               SELECT date, symbol, 'SPLIT', ratio FROM splits
               WHERE date BETWEEN ? AND ?
               ORDER BY 1, 2""", (lo, hi, lo, hi)):
        ws3.append(row)
        n_events += 1

    # ---- Symbol meta ----
    ws4 = wb.create_sheet("Symbol_Meta")
    header_row(ws4, ["Yahoo Symbol", "Long Name", "Currency", "Exchange",
                     "52W High", "52W Low", "Sessions Stored",
                     "First Date", "Last Date"],
               [14, 44, 9, 10, 11, 11, 14, 12, 12])
    n_meta = 0
    for row in conn.execute(
            """SELECT symbol, long_name, currency, exchange, week52_high,
                      week52_low, sessions, first_date, last_date
               FROM symbol_meta ORDER BY symbol"""):
        ws4.append(row)
        n_meta += 1

    # ---- Info ----
    ws5 = wb.create_sheet("Info")
    gaps = conn.execute(
        """SELECT DISTINCT s.yahoo_symbol FROM symbols s
           LEFT JOIN prices p ON p.symbol = s.yahoo_symbol
           WHERE s.yahoo_symbol IS NOT NULL AND p.symbol IS NULL""").fetchall()
    gap_txt = ", ".join(g[0] for g in gaps) or "none"
    info = [
        ("MoneyPuller price database export", label),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Price rows", n_prices),
        ("Leads in this FY", n_leads),
        ("Dividend/split events in this FY", n_events),
        ("Symbols", n_meta),
        ("Fields", "Open, High, Low, Close, Adj Close (split+div adjusted), "
                   "Volume (NSE shares)"),
        ("Note 1", "The most recent session may be partial (live intraday)."),
        ("Note 2", "Adj Close differs from Close where dividends occurred; "
                   "use Close for target-hit checks."),
        ("Note 3", "Zero-volume rows with O=H=L=C are no-trade placeholders."),
        ("Symbols absent from Yahoo (no prices)", gap_txt),
    ]
    ws5.column_dimensions["A"].width = 40
    ws5.column_dimensions["B"].width = 90
    for k, v in info:
        ws5.append([k, v])

    wb.save(out)
    return {"file": out, "prices": n_prices, "leads": n_leads,
            "events": n_events, "meta": n_meta, "secs": time.time() - t0}


def main() -> None:
    conn = sqlite3.connect(DB)
    results = []
    results.append(build(conn, FY1,
                         Path("data/MoneyPuller_Prices_FY2025-26.xlsx"),
                         "FY 2025-26 (01 Apr 2025 - 31 Mar 2026)"))
    results.append(build(conn, FY2,
                         Path("data/MoneyPuller_Prices_FY2026-27.xlsx"),
                         "FY 2026-27 to date (01 Apr 2026 - 28 Sep 2026)"))
    conn.close()
    for r in results:
        print(f"{r['file'].name}: {r['prices']:,} price rows, "
              f"{r['leads']} leads, {r['events']} corp actions, "
              f"{r['meta']} symbols ({r['secs']:.0f}s)")


if __name__ == "__main__":
    main()
