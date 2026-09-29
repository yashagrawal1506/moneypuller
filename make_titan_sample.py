"""Build a viewable Excel sample of OHLC + Volume data (TITAN.NS pilot).

Sheets:
  Titan_OHLC  - all 381 sessions with OHLC, Volume, derived change/range
                columns and a Close price line chart
  Titan_Leads - Titan's 19 leads from the master DB joined with same-day
                OHLC (CMP = previous close, per validation)
"""

import sqlite3
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DB = Path("data/moneypuller.db")
OUT = Path("data/titan_ohlc_volume_sample.xlsx")

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRICE_FMT = "#,##0.00"
VOL_FMT = "#,##0"
PCT_FMT = "0.00%"
TOP = Alignment(vertical="top")


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 20


def build():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row

    prices = conn.execute(
        """SELECT date, open, high, low, close, volume FROM prices
           WHERE symbol='TITAN.NS' ORDER BY date""").fetchall()
    leads = conn.execute(
        """SELECT a.record_date, r.cmp, r.action, r.target_1, r.target_2,
                  r.target_3, r.stop_loss, r.analyst
           FROM recommendations r JOIN articles a ON a.id = r.article_id
           WHERE r.stock_name='Titan Company' AND a.record_date IS NOT NULL
           ORDER BY a.record_date""").fetchall()
    conn.close()

    wb = Workbook()

    # ---------- Sheet 1: full OHLC + Volume ----------
    ws = wb.active
    ws.title = "Titan_OHLC"
    headers = ["Date", "Open", "High", "Low", "Close", "Volume",
               "Change %", "Range % (H-L)/C", "Gap % (O/prevC-1)"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, len(headers))
    widths = [12, 11, 11, 11, 11, 12, 10, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for r, p in enumerate(prices, start=2):
        ws.cell(row=r, column=1, value=p["date"])
        for j, key in enumerate(["open", "high", "low", "close"], start=2):
            c = ws.cell(row=r, column=j, value=p[key])
            c.number_format = PRICE_FMT
        v = ws.cell(row=r, column=6, value=p["volume"])
        v.number_format = VOL_FMT
        # derived columns as live formulas
        c = ws.cell(row=r, column=7,
                    value=f'=IFERROR(E{r}/E{r-1}-1,"")')
        c.number_format = PCT_FMT
        c = ws.cell(row=r, column=8,
                    value=f'=IFERROR((C{r}-D{r})/E{r},"")')
        c.number_format = PCT_FMT
        c = ws.cell(row=r, column=9,
                    value=f'=IFERROR(B{r}/E{r-1}-1,"")' if r > 2 else "")
        c.number_format = PCT_FMT
        for c in range(1, 10):
            ws.cell(row=r, column=c).border = BORDER

    n = len(prices) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:I{n}"

    chart = LineChart()
    chart.title = "TITAN.NS Close (Mar 2025 - Sep 2026)"
    chart.y_axis.title = "Price (INR)"
    chart.x_axis.title = "Session #"
    chart.height, chart.width = 9, 26
    data = Reference(ws, min_col=5, min_row=1, max_row=n)
    chart.add_data(data, titles_from_data=True)
    chart.legend = None
    ws.add_chart(chart, "K2")

    # ---------- Sheet 2: leads joined with prices ----------
    ws2 = wb.create_sheet("Titan_Leads")
    headers2 = ["Record Date", "Action", "CMP (prev close)",
                "Same-Day Open", "Same-Day High", "Same-Day Low",
                "Same-Day Close", "T1", "T2", "T3", "Stop-Loss",
                "T1 Hit Same Day?", "Analyst"]
    for i, h in enumerate(headers2, start=1):
        ws2.cell(row=1, column=i, value=h)
    style_header(ws2, len(headers2))
    widths2 = [12, 8, 14, 13, 13, 13, 13, 9, 9, 9, 10, 15, 38]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    price_by_date = {p["date"]: p for p in prices}
    dates_sorted = sorted(price_by_date)
    prev_close = {}
    for i, d in enumerate(dates_sorted):
        if i > 0:
            prev_close[d] = price_by_date[dates_sorted[i - 1]]["close"]

    FILL_HIT = PatternFill("solid", fgColor="C6EFCE")
    FILL_MISS = PatternFill("solid", fgColor="FCE4EC")

    for r, lead in enumerate(leads, start=2):
        d = lead["record_date"]
        p = price_by_date.get(d)
        pc = prev_close.get(d)
        ws2.cell(row=r, column=1, value=d)
        ws2.cell(row=r, column=2, value=lead["action"])
        ws2.cell(row=r, column=3, value=round(pc, 2) if pc else None
                 ).number_format = PRICE_FMT
        if p:
            for j, key in enumerate(["open", "high", "low", "close"], start=4):
                ws2.cell(row=r, column=j, value=p[key]).number_format = PRICE_FMT
        for j, key in zip(range(8, 12), ["target_1", "target_2", "target_3",
                                         "stop_loss"]):
            val = lead[key]
            if val is not None:
                ws2.cell(row=r, column=j, value=val).number_format = PRICE_FMT
        t1 = lead["target_1"]
        hit = None
        if p and t1 and lead["action"] == "Buy":
            hit = p["high"] >= t1
            cell = ws2.cell(row=r, column=12, value="YES" if hit else "no")
            cell.fill = FILL_HIT if hit else FILL_MISS
        elif p and t1 and lead["action"] == "Sell":
            hit = p["low"] <= t1
            cell = ws2.cell(row=r, column=12, value="YES" if hit else "no")
            cell.fill = FILL_HIT if hit else FILL_MISS
        ws2.cell(row=r, column=13, value=lead["analyst"])
        for c in range(1, 14):
            ws2.cell(row=r, column=c).border = BORDER
            ws2.cell(row=r, column=c).alignment = TOP

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:M{len(leads) + 1}"

    wb.save(OUT)
    print(f"wrote {OUT} with {len(prices)} price sessions "
          f"and {len(leads)} Titan leads")


if __name__ == "__main__":
    build()
